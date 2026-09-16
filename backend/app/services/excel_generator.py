import os
import logging
from typing import List, Dict, Any
import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

logging.basicConfig(level=logging.INFO, format="%(asctime)s - %(levelname)s - %(message)s")

class ExcelReportGenerator:
    @staticmethod
    def generate_excel_report(
        results: List[Dict[str, Any]], 
        failed_files: List[Dict[str, str]], 
        output_filepath: str
    ) -> str:
        """Generate formatted multi-sheet Excel workbook from candidate analysis results."""
        os.makedirs(os.path.dirname(output_filepath), exist_ok=True)
        
        with pd.ExcelWriter(output_filepath, engine='openpyxl') as writer:
            # 1. Main Sheet: Resume Analysis
            if results:
                flat_data = []
                for r in results:
                    insights = r.get("additional_insights", {}) or {}
                    scores = r.get("evaluation_scores", {}) or {}
                    
                    matched = r.get("matched_skills", [])
                    matched_str = ", ".join(matched) if isinstance(matched, list) else str(matched)
                    
                    missing = r.get("missing_skills", [])
                    missing_str = ", ".join(missing) if isinstance(missing, list) else str(missing)

                    row = {
                        "Source File": r.get("source_file", ""),
                        "Candidate Name": r.get("name", ""),
                        "Target Role": r.get("target_role", "General Role"),
                        "Match Score (%)": f"{r.get('match_score_pct', 70)}%",
                        "Recommendation": r.get("recommendation", "Shortlist"),
                        "Fit Summary": r.get("fit_summary", ""),
                        "Contact Details": r.get("contact_details", ""),
                        "University": r.get("university", ""),
                        "Degree / Course": r.get("course", ""),
                        "Discipline": r.get("discipline", ""),
                        "CGPA / %": r.get("cgpa_percentage", ""),
                        "Matched Skills": matched_str,
                        "Missing Requirements": missing_str,
                    }

                    # Add dynamic criteria columns if available
                    for crit_name, score_val in scores.items():
                        row[f"Score: {crit_name} (1-5)"] = score_val

                    row["Core Strengths"] = insights.get("technical_strength", "")
                    row["Career Potential"] = insights.get("career_potential", "")
                    row["Experience Level"] = insights.get("experience_level", "")
                    flat_data.append(row)

                df_main = pd.DataFrame(flat_data)
                df_main.to_excel(writer, sheet_name="Resume Analysis", index=False)
            else:
                pd.DataFrame({"Message": ["No successful resumes analyzed"]}).to_excel(writer, sheet_name="Resume Analysis", index=False)

            # 2. Projects Sheet
            projects_data = []
            for r in results:
                supporting = r.get("supporting_info", {}) or {}
                projects = supporting.get("projects", []) or []
                for p in projects:
                    techs = p.get("technologies", [])
                    techs_str = ", ".join(techs) if isinstance(techs, list) else str(techs)
                    projects_data.append({
                        "Candidate Name": r.get("name", ""),
                        "Project / Initiative": p.get("name", ""),
                        "Description": p.get("description", ""),
                        "Technologies / Tools": techs_str,
                        "Source File": r.get("source_file", "")
                    })
            if projects_data:
                pd.DataFrame(projects_data).to_excel(writer, sheet_name="Projects", index=False)

            # 3. Skills Analysis Sheet
            skills_data = []
            for r in results:
                skills = r.get("key_skills", [])
                skills_str = ", ".join(skills) if isinstance(skills, list) else str(skills)
                skills_data.append({
                    "Candidate Name": r.get("name", ""),
                    "Target Role": r.get("target_role", ""),
                    "Match Score": f"{r.get('match_score_pct', 70)}%",
                    "Recommendation": r.get("recommendation", ""),
                    "All Extracted Skills": skills_str,
                    "University": r.get("university", "")
                })
            if skills_data:
                pd.DataFrame(skills_data).to_excel(writer, sheet_name="Skills Analysis", index=False)

            # 4. Failed Files Sheet
            if failed_files:
                pd.DataFrame(failed_files).to_excel(writer, sheet_name="Failed Files", index=False)

        # Apply rich styling
        try:
            wb = openpyxl.load_workbook(output_filepath)
            header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
            header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
            thin_border = Border(
                left=Side(style='thin', color='CBD5E1'),
                right=Side(style='thin', color='CBD5E1'),
                top=Side(style='thin', color='CBD5E1'),
                bottom=Side(style='thin', color='CBD5E1')
            )

            for sheet_name in wb.sheetnames:
                ws = wb[sheet_name]
                for cell in ws[1]:
                    cell.fill = header_fill
                    cell.font = header_font
                    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

                ws.row_dimensions[1].height = 28

                for col in ws.columns:
                    col_letter = get_column_letter(col[0].column)
                    max_len = 0
                    for cell in col:
                        cell.border = thin_border
                        if cell.row != 1:
                            cell.alignment = Alignment(vertical="center")
                        val = str(cell.value or "")
                        max_len = max(max_len, len(val))
                    ws.column_dimensions[col_letter].width = min(max(max_len + 4, 14), 50)

            wb.save(output_filepath)
            logging.info(f"Dynamic Excel report formatted and saved successfully at {output_filepath}")
        except Exception as e:
            logging.warning(f"Styling excel file produced warning: {str(e)}")

        return output_filepath
